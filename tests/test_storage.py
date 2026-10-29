from datetime import date

from openpyxl import load_workbook

from checkin_app.storage import HEADERS, SHEET, Entry, FileLockedError, XlsxStore

import pytest


def test_roundtrip_and_upsert(tmp_path):
    store = XlsxStore(tmp_path / "sub" / "checkin.xlsx")
    assert store.load_all() == {}

    store.save(Entry(day=date(2026, 10, 20), weight=82.4, movement="Lauf", run_minutes=30,
                     goal1_name="Python", goal1_minutes=45, note="gut"))
    store.save(Entry(day=date(2026, 10, 21), movement="Spaziergang", run_minutes=99))
    # gleicher Tag -> überschreiben statt neue Zeile
    store.save(Entry(day=date(2026, 10, 20), weight=82.0, movement="Krafttraining",
                     goal1_name="Python", goal1_minutes=60))

    entries = store.load_all()
    assert len(entries) == 2
    e = entries[date(2026, 10, 20)]
    assert e.weight == 82.0 and e.movement == "Krafttraining" and e.run_minutes == 0
    assert e.goal1_minutes == 60 and e.note == ""
    # Laufminuten zählen nur bei "Lauf"
    assert entries[date(2026, 10, 21)].run_minutes == 0

    ws = load_workbook(store.path)[SHEET]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.max_row == 3
    assert ws["A2"].number_format == "DD.MM.YYYY"


def test_keeps_user_columns_and_sheets(tmp_path):
    store = XlsxStore(tmp_path / "c.xlsx")
    store.save(Entry(day=date(2026, 10, 20), weight=80))
    wb = load_workbook(store.path)
    ws = wb[SHEET]
    ws.cell(row=1, column=len(HEADERS) + 1, value="Mein Kommentar")
    ws.cell(row=2, column=len(HEADERS) + 1, value="bleibt")
    wb.create_sheet("Auswertung")["A1"] = "eigene Notizen"
    wb.save(store.path)

    store.save(Entry(day=date(2026, 10, 20), weight=81))
    wb = load_workbook(store.path)
    assert wb[SHEET].cell(row=2, column=len(HEADERS) + 1).value == "bleibt"
    assert wb["Auswertung"]["A1"].value == "eigene Notizen"
    assert store.get(date(2026, 10, 20)).weight == 81


def test_reads_user_edited_values(tmp_path):
    store = XlsxStore(tmp_path / "c.xlsx")
    store.save(Entry(day=date(2026, 10, 20)))
    wb = load_workbook(store.path)
    ws = wb[SHEET]
    ws.append(["21.10.2026", "79,5", "Lauf", "25"])
    wb.save(store.path)
    e = store.get(date(2026, 10, 21))
    assert e.weight == 79.5 and e.run_minutes == 25


def test_excel_lock_file_blocks_write(tmp_path):
    store = XlsxStore(tmp_path / "c.xlsx")
    (tmp_path / "~$c.xlsx").write_text("")
    with pytest.raises(FileLockedError):
        store.save(Entry(day=date(2026, 10, 20)))
