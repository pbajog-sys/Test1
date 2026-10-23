"""Speicherung in einer .xlsx-Datei: Blatt "Daten", eine Zeile pro Tag.

Spalten werden über die Überschrift in Zeile 1 gefunden, nicht über ihre
Position. Eigene Zusatzspalten oder weitere Blätter bleiben erhalten.
"""

from __future__ import annotations

import os
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

SHEET = "Daten"

MOVEMENTS = ("keine", "Spaziergang", "Lauf", "Krafttraining")

COL_DATE = "Datum"
COL_WEIGHT = "Gewicht (kg)"
COL_MOVEMENT = "Bewegung"
COL_RUN = "Lauf (min)"
COL_G1_NAME = "Fokusziel 1"
COL_G1_MIN = "Fokus 1 (min)"
COL_G2_NAME = "Fokusziel 2"
COL_G2_MIN = "Fokus 2 (min)"
COL_LEARN_TOTAL = "Lernzeit gesamt (min)"
COL_NOTE = "Notiz"
COL_RECORDED = "Erfasst am"

HEADERS = [
    COL_DATE, COL_WEIGHT, COL_MOVEMENT, COL_RUN,
    COL_G1_NAME, COL_G1_MIN, COL_G2_NAME, COL_G2_MIN,
    COL_LEARN_TOTAL, COL_NOTE, COL_RECORDED,
]

COLUMN_WIDTHS = {
    COL_DATE: 12, COL_WEIGHT: 13, COL_MOVEMENT: 14, COL_RUN: 11,
    COL_G1_NAME: 18, COL_G1_MIN: 13, COL_G2_NAME: 18, COL_G2_MIN: 13,
    COL_LEARN_TOTAL: 20, COL_NOTE: 40, COL_RECORDED: 17,
}


class StorageError(Exception):
    pass


class FileLockedError(StorageError):
    pass


@dataclass
class Entry:
    day: date
    weight: float | None = None
    movement: str = "keine"
    run_minutes: int = 0
    goal1_name: str = ""
    goal1_minutes: int = 0
    goal2_name: str = ""
    goal2_minutes: int = 0
    note: str = ""
    recorded_at: datetime | None = None

    @property
    def learn_minutes(self) -> int:
        return self.goal1_minutes + self.goal2_minutes

    @property
    def counted_run_minutes(self) -> int:
        return self.run_minutes if self.movement == "Lauf" else 0


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def _to_float(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _to_int(value) -> int:
    number = _to_float(value)
    return int(round(number)) if number is not None else 0


def _to_str(value) -> str:
    return "" if value is None else str(value).strip()


def _lock_file(path: Path) -> Path:
    # Excel legt beim Öffnen "~$Dateiname.xlsx" im selben Ordner an.
    return path.with_name("~$" + path.name)


class XlsxStore:
    def __init__(self, path: Path):
        self.path = Path(path)

    # -- Lesen ---------------------------------------------------------------

    def load_all(self) -> dict[date, Entry]:
        if not self.path.exists():
            return {}
        try:
            with open(self.path, "rb") as fh:
                wb = load_workbook(fh, read_only=True, data_only=True)
                try:
                    if SHEET not in wb.sheetnames:
                        return {}
                    rows = wb[SHEET].iter_rows(values_only=True)
                    header = next(rows, None)
                    if not header:
                        return {}
                    cols = {_to_str(name): i for i, name in enumerate(header) if name is not None}
                    entries: dict[date, Entry] = {}
                    for row in rows:
                        entry = self._row_to_entry(row, cols)
                        if entry is not None:
                            entries[entry.day] = entry
                    return entries
                finally:
                    wb.close()
        except PermissionError as exc:
            raise FileLockedError(str(exc)) from exc
        except StorageError:
            raise
        except Exception as exc:  # beschädigte oder fremde Datei
            raise StorageError(f"Datei konnte nicht gelesen werden: {exc}") from exc

    def get(self, day: date) -> Entry | None:
        return self.load_all().get(day)

    @staticmethod
    def _row_to_entry(row: tuple, cols: dict[str, int]) -> Entry | None:
        def cell(name):
            i = cols.get(name)
            return row[i] if i is not None and i < len(row) else None

        day = _to_date(cell(COL_DATE))
        if day is None:
            return None
        movement = _to_str(cell(COL_MOVEMENT)) or "keine"
        if movement not in MOVEMENTS:
            movement = "keine"
        recorded = cell(COL_RECORDED)
        return Entry(
            day=day,
            weight=_to_float(cell(COL_WEIGHT)),
            movement=movement,
            run_minutes=_to_int(cell(COL_RUN)),
            goal1_name=_to_str(cell(COL_G1_NAME)),
            goal1_minutes=_to_int(cell(COL_G1_MIN)),
            goal2_name=_to_str(cell(COL_G2_NAME)),
            goal2_minutes=_to_int(cell(COL_G2_MIN)),
            note=_to_str(cell(COL_NOTE)),
            recorded_at=recorded if isinstance(recorded, datetime) else None,
        )

    # -- Schreiben -----------------------------------------------------------

    def save(self, entry: Entry) -> None:
        """Legt die Zeile für entry.day an oder überschreibt sie."""
        if _lock_file(self.path).exists():
            raise FileLockedError("Die Datei ist gerade in Excel geöffnet.")
        self.path.parent.mkdir(parents=True, exist_ok=True)

        try:
            wb = load_workbook(self.path) if self.path.exists() else self._new_workbook()
        except PermissionError as exc:
            raise FileLockedError(str(exc)) from exc
        except Exception as exc:
            raise StorageError(f"Datei konnte nicht gelesen werden: {exc}") from exc

        ws = wb[SHEET] if SHEET in wb.sheetnames else self._add_sheet(wb)
        cols = self._ensure_headers(ws)

        target_row = None
        for r in range(2, ws.max_row + 1):
            if _to_date(ws.cell(row=r, column=cols[COL_DATE]).value) == entry.day:
                target_row = r
                break
        if target_row is None:
            target_row = self._first_empty_row(ws, cols[COL_DATE])

        values = {
            COL_DATE: entry.day,
            COL_WEIGHT: entry.weight,
            COL_MOVEMENT: entry.movement,
            COL_RUN: entry.counted_run_minutes or None,
            COL_G1_NAME: entry.goal1_name or None,
            COL_G1_MIN: entry.goal1_minutes if entry.goal1_name else None,
            COL_G2_NAME: entry.goal2_name or None,
            COL_G2_MIN: entry.goal2_minutes if entry.goal2_name else None,
            COL_LEARN_TOTAL: entry.learn_minutes,
            COL_NOTE: entry.note or None,
            COL_RECORDED: (entry.recorded_at or datetime.now()).replace(microsecond=0),
        }
        formats = {COL_DATE: "DD.MM.YYYY", COL_WEIGHT: "0.0", COL_RECORDED: "DD.MM.YYYY HH:MM"}
        for name, value in values.items():
            c = ws.cell(row=target_row, column=cols[name])
            c.value = value  # ws.cell(value=None) würde alte Werte stehen lassen
            if name in formats:
                c.number_format = formats[name]

        self._write_atomically(wb)

    def _write_atomically(self, wb: Workbook) -> None:
        tmp = self.path.with_name(self.path.stem + ".speichern.tmp")
        try:
            wb.save(tmp)
            if self.path.exists():
                shutil.copy2(self.path, self.path.with_name(self.path.name + ".bak"))
            os.replace(tmp, self.path)
        except PermissionError as exc:
            raise FileLockedError(str(exc)) from exc
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    @staticmethod
    def _new_workbook() -> Workbook:
        wb = Workbook()
        wb.remove(wb.active)
        return wb

    @staticmethod
    def _add_sheet(wb: Workbook):
        ws = wb.create_sheet(SHEET, 0)
        ws.freeze_panes = "A2"
        return ws

    @staticmethod
    def _ensure_headers(ws) -> dict[str, int]:
        cols: dict[str, int] = {}
        last = 0
        for c in range(1, ws.max_column + 1):
            name = _to_str(ws.cell(row=1, column=c).value)
            if name:
                cols.setdefault(name, c)
                last = c
        for name in HEADERS:
            if name not in cols:
                last += 1
                cell = ws.cell(row=1, column=last, value=name)
                cell.font = Font(bold=True)
                ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS.get(name, 14)
                cols[name] = last
        return cols

    @staticmethod
    def _first_empty_row(ws, date_col: int) -> int:
        r = ws.max_row
        while r >= 2 and ws.cell(row=r, column=date_col).value in (None, ""):
            r -= 1
        return r + 1
